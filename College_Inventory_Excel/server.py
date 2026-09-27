"""Local inventory web server. Excel workbook is the persistent data store."""
import json
import mimetypes
import threading
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent
BOOK = ROOT / 'data' / 'inventory.xlsx'
FRONT = ROOT / 'frontend'
HEADERS = ['id', 'name', 'category', 'quantity', 'location', 'updated_at']
LOCK = threading.RLock()


def ensure_book():
    if not BOOK.exists():
        BOOK.parent.mkdir(exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = 'Items'
        ws.append(HEADERS)
        ws.freeze_panes = 'A2'
        for col, width in {'A': 10, 'B': 30, 'C': 22, 'D': 13, 'E': 25, 'F': 24}.items():
            ws.column_dimensions[col].width = width
        wb.save(BOOK)


def read_items():
    ensure_book()
    wb = load_workbook(BOOK, read_only=True, data_only=True)
    try:
        return [dict(zip(HEADERS, row)) for row in list(wb.active.values)[1:] if row[0] is not None]
    finally:
        wb.close()


def validate(payload):
    if not isinstance(payload, dict):
        raise ValueError('Enter item details.')
    name = str(payload.get('name', '')).strip()
    category = str(payload.get('category', '')).strip()
    location = str(payload.get('location', '')).strip()
    try:
        quantity = int(payload.get('quantity', ''))
    except (ValueError, TypeError):
        raise ValueError('Quantity must be a whole number.')
    if not name or not category or not location or any(len(x) > 100 for x in (name, category, location)):
        raise ValueError('Name, category and location are required (maximum 100 characters each).')
    if quantity < 0 or quantity > 1000000:
        raise ValueError('Quantity must be between 0 and 1,000,000.')
    return name, category, quantity, location


class Handler(BaseHTTPRequestHandler):
    def reply(self, status, value):
        data = json.dumps(value).encode('utf-8')
        self.send_response(status)
        self.send_header('Content-Type', 'application/json; charset=utf-8')
        self.send_header('Content-Length', str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def body(self):
        size = int(self.headers.get('Content-Length', '0'))
        if size < 1 or size > 10000:
            raise ValueError('Request must contain item details under 10 KB.')
        return json.loads(self.rfile.read(size))

    def do_GET(self):
        path = urlsplit(self.path).path
        if path == '/api/items':
            try:
                with LOCK:
                    self.reply(200, read_items())
            except (OSError, PermissionError):
                self.reply(503, {'error': 'Close inventory.xlsx in Excel and try again.'})
            return
        name = 'index.html' if path == '/' else path.removeprefix('/assets/') if path.startswith('/assets/') else ''
        target = FRONT / name
        if not name or name not in ('index.html', 'style.css', 'app.js') or not target.is_file():
            self.send_error(404)
            return
        data = target.read_bytes()
        self.send_response(200)
        self.send_header('Content-Type', mimetypes.guess_type(target.name)[0] or 'text/plain')
        self.send_header('Content-Length', str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def mutate(self, mode, item_id=None):
        try:
            payload = self.body() if mode != 'delete' else None
            fields = validate(payload) if mode != 'delete' else None
            with LOCK:
                ensure_book()
                wb = load_workbook(BOOK)
                ws = wb.active
                rows = list(ws.iter_rows(min_row=2))
                row = next((cells for cells in rows if cells[0].value == item_id), None)
                if mode == 'add':
                    new_id = max([cells[0].value for cells in rows] or [0]) + 1
                    ws.append([new_id, *fields, datetime.now().strftime('%Y-%m-%d %H:%M')])
                elif row is None:
                    wb.close()
                    return self.reply(404, {'error': 'Item not found.'})
                elif mode == 'edit':
                    for cell, value in zip(row[1:], [*fields, datetime.now().strftime('%Y-%m-%d %H:%M')]):
                        cell.value = value
                else:
                    ws.delete_rows(row[0].row)
                wb.save(BOOK)
                wb.close()
            self.reply(201 if mode == 'add' else 200, {'ok': True})
        except (ValueError, json.JSONDecodeError) as error:
            self.reply(400, {'error': str(error)})
        except (OSError, PermissionError):
            self.reply(503, {'error': 'Cannot save workbook. Close inventory.xlsx in Excel and try again.'})

    def item_id(self):
        try:
            return int(urlsplit(self.path).path.split('/')[-1])
        except ValueError:
            return None

    def do_POST(self):
        if urlsplit(self.path).path == '/api/items':
            self.mutate('add')
        else:
            self.reply(404, {'error': 'Route not found.'})

    def do_PUT(self):
        item_id = self.item_id()
        if urlsplit(self.path).path.startswith('/api/items/') and item_id is not None:
            self.mutate('edit', item_id)
        else:
            self.reply(404, {'error': 'Route not found.'})

    def do_DELETE(self):
        item_id = self.item_id()
        if urlsplit(self.path).path.startswith('/api/items/') and item_id is not None:
            self.mutate('delete', item_id)
        else:
            self.reply(404, {'error': 'Route not found.'})


if __name__ == '__main__':
    ensure_book()
    print('Open http://127.0.0.1:8000 in your browser')
    ThreadingHTTPServer(('127.0.0.1', 8000), Handler).serve_forever()
